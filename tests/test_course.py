"""Checks for safe course ZIP import and source-linked retrieval."""

import io
import tempfile
import unittest
import zipfile
from pathlib import Path

from oslab.course import CourseSearch, import_course


class CourseIndexTests(unittest.TestCase):
    def test_local_import_preserves_source_location_and_skips_planning_file(self):
        import pymupdf
        from pptx import Presentation

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            pdf = pymupdf.open()
            page = pdf.new_page()
            page.insert_text((72, 72), "A page fault occurs when a process accesses a virtual page that is not in memory. " * 2)
            pdf_bytes = pdf.tobytes()
            pdf.close()

            slides = Presentation()
            slide = slides.slides.add_slide(slides.slide_layouts[1])
            slide.shapes.title.text = "Semaphores"
            slide.placeholders[1].text = "A semaphore coordinates threads using wait and signal operations. " * 2
            stream = io.BytesIO()
            slides.save(stream)

            archive = root / "course.zip"
            with zipfile.ZipFile(archive, "w") as output:
                output.writestr("OS/Theory/paging.pdf", pdf_bytes)
                output.writestr("OS/Theory/synchronization.pptx", stream.getvalue())
                output.writestr("OS/Problem Statement LLM.docx", b"project instructions are not course data")
                output.writestr("../outside.pdf", pdf_bytes)
            index_path = root / "private" / "index.json"
            payload = import_course(archive, index_path)
            self.assertEqual(len(payload["documents"]), 2)
            self.assertEqual(len(payload["skipped"]), 2)
            self.assertTrue(index_path.is_file())
            index = CourseSearch.load(index_path)
            self.assertEqual(index.search("page fault")[0]["location"], "page 1")
            self.assertEqual(index.search("semaphore threads")[0]["location"], "slide 1")
            self.assertEqual(index.status()["passages"], len(payload["records"]))
            enriched = CourseSearch.load(index_path, include_demo=True)
            self.assertEqual(enriched.status()["supplemental_passages"], 21)
            self.assertEqual(enriched.status()["passages"], len(payload["records"]))
            self.assertTrue(any(item["group"] == "Project note" for item in
                                enriched.search("pipe reader EOF close", limit=5)))


if __name__ == "__main__":
    unittest.main()
