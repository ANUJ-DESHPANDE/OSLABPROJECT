"""Local, private course document ingestion and cited passage search."""

import hashlib
import io
import json
import math
import re
import unicodedata
import zipfile
from collections import Counter
from pathlib import Path, PurePosixPath
from xml.etree import ElementTree

from sklearn.decomposition import TruncatedSVD
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

from .knowledge import section_chunks, tokens

INDEX_PATH = Path("data/private/processed/course_index.json")
REFERENCE_INDEX_PATH = Path("data/private/processed/textbook_resources_index.json")
SUPPORTED = {".pdf", ".pptx", ".docx"}
MAX_ARCHIVE_SIZE = 200_000_000
MAX_ENTRY_SIZE = 20_000_000


def _clean(value):
    """Normalize copy/paste quirks while preserving paragraph boundaries."""
    value = unicodedata.normalize("NFKC", value)
    return re.sub(r"[ \t]+", " ", value).strip()


def _pages(name, data):
    """Yield (location, text) so every answer can point to a page or slide."""
    extension = PurePosixPath(name).suffix.lower()
    if extension == ".pdf":
        import pymupdf

        with pymupdf.open(stream=data, filetype="pdf") as document:
            for number, page in enumerate(document, 1):
                yield f"page {number}", page.get_text("text")
    elif extension == ".pptx":
        from pptx import Presentation

        presentation = Presentation(io.BytesIO(data))
        for number, slide in enumerate(presentation.slides, 1):
            yield f"slide {number}", "\n".join(
                shape.text for shape in slide.shapes if shape.has_text_frame
            )
    elif extension == ".docx":
        with zipfile.ZipFile(io.BytesIO(data)) as document:
            root = ElementTree.fromstring(document.read("word/document.xml"))
        namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        paragraphs = []
        for paragraph in root.iter(namespace + "p"):
            line = "".join(node.text or "" for node in paragraph.iter(namespace + "t"))
            if line.strip():
                paragraphs.append(line)
        yield "document", "\n\n".join(paragraphs)


def import_course(archive_path, index_path=INDEX_PATH, *, source_group=None, max_entry_size=MAX_ENTRY_SIZE):
    """Turn a source ZIP into local passage records; never publish source files.

    Reject unsafe/oversized entries, deduplicate identical files by SHA-256, then
    store each readable chunk with its document and page/slide location.
    """
    archive_path, index_path = Path(archive_path), Path(index_path)
    if not archive_path.is_file():
        raise ValueError(f"Course ZIP not found: {archive_path}")
    records, documents, skipped, seen_hashes = [], [], [], set()
    with zipfile.ZipFile(archive_path) as archive:
        entries = [entry for entry in archive.infolist() if not entry.is_dir()]
        if len(entries) > 500 or sum(entry.file_size for entry in entries) > MAX_ARCHIVE_SIZE:
            raise ValueError("Course archive exceeds the allowed size")
        for entry in entries:
            name = entry.filename.replace("\\", "/")
            path = PurePosixPath(name)
            if path.is_absolute() or ".." in path.parts or ":" in name:
                skipped.append({"source": name, "reason": "unsafe archive path"})
                continue
            suffix = path.suffix.lower()
            if entry.file_size > max_entry_size:
                skipped.append({"source": name, "reason": "file too large"})
                continue
            if path.name.startswith("~$") or path.name.lower() == "problem statement llm.docx":
                skipped.append({"source": name, "reason": "temporary or project planning file"})
                continue
            if suffix not in SUPPORTED:
                skipped.append({"source": name, "reason": "format not text-extractable here"})
                continue
            try:
                data = archive.read(entry)
                digest = hashlib.sha256(data).hexdigest()
                if digest in seen_hashes:
                    skipped.append({"source": name, "reason": "duplicate file"})
                    continue
                seen_hashes.add(digest)
                before = len(records)
                if source_group:
                    group = source_group
                elif name.startswith("Textbook 8th/"):
                    group = "OS Concepts textbook (8th ed)"
                elif name.startswith("Textbook companion/"):
                    group = "OS Concepts companion slides (10th ed)"
                else:
                    group = "Lab" if "/Lab/" in f"/{name}/" else "Theory"
                for location, raw_text in _pages(name, data):
                    content = _clean(raw_text)
                    for part, chunk in enumerate(section_chunks(content, max_words=210), 1):
                        chunk = _clean(chunk)
                        if len(chunk) < 60:
                            continue
                        record_id = hashlib.sha256(f"{name}|{location}|{part}|{chunk}".encode()).hexdigest()[:16]
                        records.append({"id": record_id, "source": name, "title": path.stem,
                                        "group": group, "format": suffix[1:], "location": location,
                                        "part": part, "text": chunk})
                documents.append({"source": name, "format": suffix[1:], "group": group,
                                  "chunks": len(records) - before,
                                  "sha256": digest})
            except Exception as error:
                skipped.append({"source": name, "reason": f"extraction failed: {type(error).__name__}"})
    if not records:
        raise ValueError("No searchable text could be extracted from this archive")
    payload = {"schema_version": 1, "documents": documents, "skipped": skipped, "records": records}
    index_path.parent.mkdir(parents=True, exist_ok=True)
    index_path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    return payload


class CourseSearch:
    """Rank local passages while retaining their original source locations.

    BM25 uses term counts; LSA compresses TF-IDF vectors. Both are fitted over
    local passages when the app starts. No model training or internet call occurs.
    """

    def __init__(self, payload):
        self.payload = payload
        self.records = payload["records"]
        self.course_count = payload.get("course_passages", len(self.records))
        self.reference_count = payload.get("reference_passages", 0)
        self.docs = [tokens(record["text"] + " " + record["title"]) for record in self.records]
        self.counts = [Counter(doc) for doc in self.docs]
        self.df = Counter(term for doc in self.docs for term in set(doc))
        self.avg_length = sum(map(len, self.docs)) / max(1, len(self.docs))
        self.vectorizer = TfidfVectorizer(stop_words="english", ngram_range=(1, 2), sublinear_tf=True)
        matrix = self.vectorizer.fit_transform(record["text"] + " " + record["title"] for record in self.records)
        dimensions = min(48, matrix.shape[0] - 1, matrix.shape[1] - 1)
        self.svd = TruncatedSVD(n_components=dimensions, random_state=7) if dimensions >= 2 else None
        self.dense = self.svd.fit_transform(matrix) if self.svd else matrix.toarray()

    @classmethod
    def load(cls, index_path=INDEX_PATH, include_demo=False, reference_index_path=None):
        path = Path(index_path)
        if not path.exists():
            return None
        payload = json.loads(path.read_text(encoding="utf-8"))
        payload["course_passages"] = len(payload["records"])
        if reference_index_path is not None and Path(reference_index_path).exists():
            reference = json.loads(Path(reference_index_path).read_text(encoding="utf-8"))
            payload["reference_passages"] = len(reference["records"])
            payload["reference_documents"] = len(reference["documents"])
            payload["reference_inventory"] = reference["documents"]
            payload["reference_skipped"] = reference["skipped"]
            payload["reference_formats"] = dict(Counter(item["format"] for item in reference["documents"]))
            payload["reference_groups"] = dict(Counter(item["group"] for item in reference["documents"]))
            payload["records"].extend(reference["records"])
        if include_demo:
            from .knowledge import load_demo

            payload["records"].extend({
                "id": item["id"], "source": "Project lab note: " + item["source"],
                "title": item["experiment_name"], "group": "Project note",
                "location": item["page_section"], "text": item["text"],
            } for item in load_demo())
        return cls(payload)

    def status(self):
        return {"ready": True, "documents": len(self.payload["documents"]),
                "passages": self.course_count,
                "reference_documents": self.payload.get("reference_documents", 0),
                "reference_passages": self.reference_count,
                "reference_inventory": self.payload.get("reference_inventory", []),
                "reference_skipped": self.payload.get("reference_skipped", []),
                "reference_formats": self.payload.get("reference_formats", {}),
                "reference_groups": self.payload.get("reference_groups", {}),
                "supplemental_passages": len(self.records) - self.course_count - self.reference_count,
                "skipped": len(self.payload["skipped"]),
                "formats": dict(Counter(item["format"] for item in self.payload["documents"])),
                "skip_reasons": dict(Counter(item["reason"] for item in self.payload["skipped"])),
                "source_inventory": self.payload["documents"],
                "skipped_entries": self.payload["skipped"]}

    def search(self, question, limit=5, method="hybrid", scope="course"):
        """Return ranked source passages, optionally filtering course vs book.

        A retrieval hit is evidence for an answer, not a generated answer itself.
        """
        if method not in {"bm25", "lsa", "hybrid"}:
            raise ValueError("Search method must be bm25, lsa, or hybrid")
        if scope not in {"course", "all", "textbook"}:
            raise ValueError("Search scope must be course, all, or textbook")
        terms = set(tokens(question))
        if not terms:
            return []
        total = len(self.docs)
        lexical = []  # BM25 rewards matching terms but corrects for passage length.
        exact_phrase = re.sub(r"[?!.\s]+$", "", question).lower()
        for index, (document, counts) in enumerate(zip(self.docs, self.counts)):
            score = sum(
                math.log(1 + (total - self.df[term] + .5) / (self.df[term] + .5))
                * counts[term] * 2.2 / (counts[term] + 1.2 * (.25 + .75 * len(document) / self.avg_length))
                for term in terms if counts[term]
            )
            if len(exact_phrase.split()) >= 3 and exact_phrase in self.records[index]["text"].lower():
                score += 6
            lexical.append(score)
        query_vector = self.vectorizer.transform([question])
        dense_query = self.svd.transform(query_vector) if self.svd else query_vector.toarray()
        # LSA compares compressed local text vectors; it is not neural embeddings.
        semantic = cosine_similarity(dense_query, self.dense)[0].tolist()
        lexical_order = sorted(range(total), key=lambda i: (-lexical[i], i))
        semantic_order = sorted(range(total), key=lambda i: (-semantic[i], i))
        lexical_rank = {index: rank for rank, index in enumerate(lexical_order, 1)}
        semantic_rank = {index: rank for rank, index in enumerate(semantic_order, 1)}
        ranked = []
        for index in range(total):
            if lexical[index] <= 0 and semantic[index] < 0.08:
                continue
            # Reciprocal rank fusion lets either method contribute a strong hit.
            score = {"bm25": lexical[index], "lsa": semantic[index],
                     "hybrid": 2 / (60 + lexical_rank[index]) + 1 / (60 + semantic_rank[index])}[method]
            if method == "hybrid" and len(exact_phrase.split()) >= 3 and exact_phrase in self.records[index]["text"].lower():
                score += .02
            ranked.append((score, index))
        ranked.sort(key=lambda item: (-item[0], item[1]))
        # Keep the first result per document/location so one long page does not crowd out others.
        results, seen = [], set()
        for score, index in ranked:
            record = self.records[index]
            is_textbook = record["group"].startswith("OS Concepts")
            if scope == "course" and is_textbook or scope == "textbook" and not is_textbook:
                continue
            key = (record["source"], record["location"])
            if key in seen:
                continue
            seen.add(key)
            results.append({"source": record["source"], "title": record["title"],
                            "location": record["location"], "group": record["group"],
                            "text": record["text"], "score": round(score, 5),
                            "bm25_score": round(lexical[index], 3),
                            "lsa_score": round(semantic[index], 3), "method": method})
            if len(results) == limit:
                break
        return results
